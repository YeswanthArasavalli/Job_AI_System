"""
Resume File Parser -- Extracts text from PDF, DOCX, and JSON resume files.

Supports:
  - PDF (via pdfplumber)
  - DOCX (via python-docx)
  - JSON (native)
  - Raw text fallback
"""

import json
import os
import re
from typing import Optional

from utils.logger import setup_logger

logger = setup_logger(__name__)


class ResumeFileParser:
    """
    Parse resume files into structured text or JSON.

    Usage:
        parser = ResumeFileParser()
        text = parser.extract_text("resume.pdf")
        data = parser.parse_to_dict("resume.pdf")
    """

    def extract_text(self, file_path: str = None, file_bytes: bytes = None, filename: str = "") -> str:
        """
        Extract raw text from a resume file.

        Args:
            file_path: Path to file on disk.
            file_bytes: Raw bytes (for uploaded files).
            filename: Original filename (for type detection).

        Returns:
            Extracted text string.
        """
        ext = ""
        if file_path:
            ext = os.path.splitext(file_path)[1].lower()
            filename = os.path.basename(file_path)
        elif filename:
            ext = os.path.splitext(filename)[1].lower()

        logger.info(f"Parsing resume: {filename} (type: {ext or 'unknown'})")

        try:
            if ext == ".pdf":
                return self._parse_pdf(file_path, file_bytes)
            elif ext == ".docx":
                return self._parse_docx(file_path, file_bytes)
            elif ext == ".json":
                return self._parse_json(file_path, file_bytes)
            elif ext == ".txt":
                if file_bytes:
                    return file_bytes.decode("utf-8", errors="replace")
                return open(file_path, "r", encoding="utf-8").read()
            else:
                # Try as text
                if file_bytes:
                    return file_bytes.decode("utf-8", errors="replace")
                logger.warning(f"Unknown file type: {ext}. Trying as text.")
                return open(file_path, "r", encoding="utf-8").read()
        except Exception as e:
            logger.error(f"Resume parsing failed: {e}")
            raise ValueError(f"Could not parse resume file: {e}")

    def parse_to_dict(self, file_path: str = None, file_bytes: bytes = None, filename: str = "") -> Optional[dict]:
        """
        Parse resume to dict. Only works for JSON files.
        For PDF/DOCX, returns None (use extract_text + LLM instead).
        """
        ext = ""
        if file_path:
            ext = os.path.splitext(file_path)[1].lower()
        elif filename:
            ext = os.path.splitext(filename)[1].lower()

        if ext == ".json":
            raw = self._parse_json(file_path, file_bytes)
            try:
                return json.loads(raw)
            except json.JSONDecodeError:
                return None
        return None

    # ================================================================
    # Parsers
    # ================================================================

    def _parse_pdf(self, file_path: str = None, file_bytes: bytes = None) -> str:
        """Extract text from PDF using pdfplumber."""
        try:
            import pdfplumber
        except ImportError:
            raise ImportError(
                "pdfplumber is required for PDF parsing. "
                "Install with: pip install pdfplumber"
            )

        pages_text = []
        if file_path:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        pages_text.append(text)
        elif file_bytes:
            import io
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        pages_text.append(text)

        full_text = "\n\n".join(pages_text)
        if not full_text.strip():
            raise ValueError("PDF contains no extractable text (may be image-based)")

        logger.info(f"PDF parsed: {len(pages_text)} pages, {len(full_text)} chars")
        return full_text

    def _parse_docx(self, file_path: str = None, file_bytes: bytes = None) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            from docx import Document
        except ImportError:
            raise ImportError(
                "python-docx is required for DOCX parsing. "
                "Install with: pip install python-docx"
            )

        if file_path:
            doc = Document(file_path)
        elif file_bytes:
            import io
            doc = Document(io.BytesIO(file_bytes))
        else:
            raise ValueError("No file provided")

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)

        if not full_text.strip():
            raise ValueError("DOCX contains no text")

        logger.info(f"DOCX parsed: {len(paragraphs)} paragraphs, {len(full_text)} chars")
        return full_text

    def _parse_json(self, file_path: str = None, file_bytes: bytes = None) -> str:
        """Read JSON file as raw text."""
        if file_bytes:
            return file_bytes.decode("utf-8", errors="replace")
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
